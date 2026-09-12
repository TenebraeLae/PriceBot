import asyncio
from datetime import datetime, timedelta
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from sqlalchemy import select, update

from pricebot.db.models import Order, PromoCode
from pricebot.domain.errors import DomainError
from pricebot.domain.payments import OrderStatus
from pricebot.domain.status import apply_admin_status
from pricebot.domain.telegram import call_with_retry
from pricebot.services.backup import rotate_backups, write_backup_file
from pricebot.services.notify import get_notification_log, reset_notifications
from pricebot.services.load_search import run_search_load
from tests.conftest import init_data_headers
from tests.test_miniapp_api import _cleanup, _client, _import_catalog

MOSCOW = ZoneInfo("Europe/Moscow")
ADMIN = {"X-Telegram-Id": "1001"}
USER = 555


def _seed(client, tmp_path: Path) -> None:
    _import_catalog(
        client,
        tmp_path,
        [["CEM-M500", "Цемент М500", "450.00", 12, "Сухие смеси", "", 10]],
    )


def _checkout(client, settings, extra=None):
    web = init_data_headers(settings, USER)
    client.post("/api/v1/cart/items", headers=web, json={"sku": "CEM-M500", "qty": 1})
    return client.post("/api/v1/checkout", headers=web, json=extra or {})


def test_admin_forbidden_and_order_filters(tmp_path: Path) -> None:
    client, settings, factory, engine = _client(tmp_path)
    reset_notifications()
    try:
        _seed(client, tmp_path)
        created = _checkout(client, settings)
        number = created.json()["number"]
        denied = client.get("/api/v1/admin/orders", headers={"X-Telegram-Id": "9"})
        assert denied.status_code == 403
        found = client.get(
            "/api/v1/admin/orders",
            params={"number": number, "telegram_id": USER, "sku": "CEM-M500", "status": "awaiting_payment"},
            headers=ADMIN,
        )
        assert found.status_code == 200
        assert found.json()["items"][0]["number"] == number
    finally:
        _cleanup(engine)


def test_admin_status_transitions_and_cannot_set_paid(tmp_path: Path) -> None:
    client, settings, factory, engine = _client(tmp_path)
    reset_notifications()
    try:
        _seed(client, tmp_path)
        number = _checkout(client, settings).json()["number"]
        paid = client.post(
            f"/api/v1/admin/orders/{number}/status",
            json={"status": "paid"},
            headers=ADMIN,
        )
        assert paid.status_code == 400

        cancelled = client.post(
            f"/api/v1/admin/orders/{number}/status",
            json={"status": "cancelled"},
            headers=ADMIN,
        )
        assert cancelled.status_code == 200
        assert cancelled.json()["status"] == "cancelled"
        assert get_notification_log().count(kind="order_cancelled", order_number=number) == 1

        other = _checkout(client, settings).json()["number"]

        async def _mark_paid() -> None:
            async with factory() as session:
                await session.execute(
                    update(Order).where(Order.number == other).values(status="paid")
                )
                await session.commit()

        asyncio.run(_mark_paid())
        processing = client.post(
            f"/api/v1/admin/orders/{other}/status",
            json={"status": "processing"},
            headers=ADMIN,
        )
        assert processing.json()["status"] == "processing"
        done = client.post(
            f"/api/v1/admin/orders/{other}/status",
            json={"status": "done"},
            headers=ADMIN,
        )
        assert done.json()["status"] == "done"
        assert get_notification_log().count(kind="order_done", order_number=other) == 1
    finally:
        _cleanup(engine)


def test_admin_users_block_unblock(tmp_path: Path) -> None:
    client, settings, factory, engine = _client(tmp_path)
    try:
        _seed(client, tmp_path)
        web = init_data_headers(settings, USER)
        client.get("/api/v1/search", params={"q": "цемент"}, headers=web)
        listed = client.get("/api/v1/admin/users", headers=ADMIN)
        assert listed.status_code == 200
        ids = [item["telegram_id"] for item in listed.json()["items"]]
        assert USER in ids
        blocked = client.post(f"/api/v1/admin/users/{USER}/block", headers=ADMIN)
        assert blocked.json()["status"] == "blocked"
        search = client.get("/api/v1/search", params={"q": "цемент"}, headers=web)
        assert search.json()["reason"] == "user_blocked"
        client.post("/api/v1/cart/items", headers=web, json={"sku": "CEM-M500", "qty": 1})
        refused = client.post("/api/v1/checkout", headers=web, json={})
        assert refused.status_code == 403
        client.post(f"/api/v1/admin/users/{USER}/unblock", headers=ADMIN)
    finally:
        _cleanup(engine)


def test_promo_create_checkout_and_expired(tmp_path: Path) -> None:
    client, settings, factory, engine = _client(tmp_path)
    try:
        _seed(client, tmp_path)
        now = datetime(2026, 9, 12, 12, 0, tzinfo=MOSCOW)
        created = client.post(
            "/api/v1/admin/promos",
            json={
                "code": "sale10",
                "kind": "percent",
                "value": "10",
                "starts_at": (now - timedelta(days=1)).isoformat(),
                "ends_at": (now + timedelta(days=1)).isoformat(),
                "max_uses": 5,
                "active": True,
            },
            headers=ADMIN,
        )
        assert created.status_code == 200
        assert created.json()["code"] == "SALE10"
        applied = _checkout(client, settings, extra={"promo_code": "SALE10"})
        assert applied.status_code == 200
        assert applied.json()["amount"] == "405.00"
        assert applied.json()["promo_code"] == "SALE10"

        async def _used() -> int:
            async with factory() as session:
                row = (await session.execute(select(PromoCode))).scalar_one()
                return int(row.used_count)

        assert asyncio.run(_used()) == 1

        expired = client.post(
            "/api/v1/admin/promos",
            json={
                "code": "OLD",
                "kind": "fixed",
                "value": "50",
                "starts_at": (now - timedelta(days=10)).isoformat(),
                "ends_at": (now - timedelta(days=1)).isoformat(),
                "active": True,
            },
            headers=ADMIN,
        )
        assert expired.status_code == 200
        client.post("/api/v1/admin/promos/OLD/disable", headers=ADMIN)
        refused = _checkout(client, settings, extra={"promo_code": "OLD"})
        assert refused.status_code == 400
    finally:
        _cleanup(engine)


def test_ticket_create_and_admin_reply(tmp_path: Path) -> None:
    client, settings, factory, engine = _client(tmp_path)
    reset_notifications()
    try:
        web = init_data_headers(settings, USER)
        created = client.post("/api/v1/tickets", headers=web, json={"text": "Нет цемента"})
        assert created.status_code == 200
        ticket = created.json()["ticket"]
        assert created.json()["status"] == "open"
        opened = client.get("/api/v1/admin/tickets", headers=ADMIN)
        assert opened.json()["items"][0]["ticket"] == ticket
        reply = client.post(
            f"/api/v1/admin/tickets/{ticket}/reply",
            json={"text": "Завезём завтра"},
            headers=ADMIN,
        )
        assert reply.json()["status"] == "closed"
        assert reply.json()["admin_reply"] == "Завезём завтра"
        assert get_notification_log().count(kind="ticket_reply") == 1
        all_tickets = client.get("/api/v1/admin/tickets", params={"status": "all"}, headers=ADMIN)
        assert all_tickets.json()["items"]
    finally:
        _cleanup(engine)


def test_catalog_xlsx_export(tmp_path: Path) -> None:
    client, settings, factory, engine = _client(tmp_path)
    try:
        _seed(client, tmp_path)
        denied = client.get("/api/v1/admin/catalog.xlsx", headers={"X-Telegram-Id": "9"})
        assert denied.status_code == 403
        exported = client.get("/api/v1/admin/catalog.xlsx", headers=ADMIN)
        assert exported.status_code == 200
        book = load_workbook(BytesIO(exported.content), read_only=True)
        rows = list(book.active.iter_rows(values_only=True))
        book.close()
        assert list(rows[0]) == [
            "sku",
            "name",
            "category",
            "description",
            "price",
            "stock",
            "unit",
            "photo",
            "active",
            "sort",
        ]
        skus = [row[0] for row in rows[1:]]
        assert "CEM-M500" in skus
    finally:
        _cleanup(engine)


def test_telegram_retry_respects_retry_after() -> None:
    slept: list[float] = []

    class FakeRetry(Exception):
        def __init__(self, retry_after: float) -> None:
            self.retry_after = retry_after

    calls = {"n": 0}

    async def _op():
        calls["n"] += 1
        if calls["n"] == 1:
            raise FakeRetry(1.5)
        return "ok"

    async def _sleep(delay: float) -> None:
        slept.append(delay)

    assert asyncio.run(call_with_retry(_op, sleeper=_sleep)) == "ok"
    assert slept == [1.5]


def test_backup_rotation(tmp_path: Path) -> None:
    now = datetime(2026, 9, 12, 12, 0, tzinfo=MOSCOW)
    fresh = write_backup_file(tmp_path, b"new", now=now)
    old = tmp_path / "pricebot-old.sql"
    old.write_bytes(b"old")
    old_mtime = (now - timedelta(days=10)).timestamp()
    import os

    os.utime(old, (old_mtime, old_mtime))
    removed = rotate_backups(tmp_path, retention_days=7, now=now)
    assert old in removed
    assert fresh.exists()
    assert not old.exists()


def test_telegram_webhook_bad_secret_403(tmp_path: Path) -> None:
    client, settings, factory, engine = _client(tmp_path)
    try:
        bad = client.post(
            "/telegram/webhook",
            json={"update_id": 1},
            headers={"X-Telegram-Bot-Api-Secret-Token": "wrong"},
        )
        assert bad.status_code == 403
        good = client.post(
            "/telegram/webhook",
            json={"update_id": 1},
            headers={"X-Telegram-Bot-Api-Secret-Token": settings.telegram_webhook_secret},
        )
        assert good.status_code == 200
    finally:
        _cleanup(engine)


def test_load_search_n100_no_5xx(tmp_path: Path) -> None:
    client, settings, factory, engine = _client(tmp_path)
    try:
        _seed(client, tmp_path)
        headers = init_data_headers(settings, USER)
        report = run_search_load(client, headers, n=100, q="цемент")
        assert report["errors_5xx"] == 0
        assert report["n"] == 100
        assert report["concurrent"] == 100
        assert report["p50_ms"] >= 0
        assert report["p95_ms"] >= report["p50_ms"]
    finally:
        _cleanup(engine)


def test_admin_status_domain_rules() -> None:
    assert apply_admin_status("awaiting_payment", "cancelled") is OrderStatus.cancelled
    assert apply_admin_status("paid", "processing") is OrderStatus.processing
    assert apply_admin_status("processing", "done") is OrderStatus.done
    assert apply_admin_status("paid", "error") is OrderStatus.error
    try:
        apply_admin_status("awaiting_payment", "paid")
        raise AssertionError("paid must be refused")
    except DomainError as exc:
        assert exc.code == "paid_not_allowed"
    try:
        apply_admin_status("done", "error")
        raise AssertionError("done cannot go to error")
    except DomainError as exc:
        assert exc.code == "invalid_transition"


def test_miniapp_has_promo_field() -> None:
    from pricebot.web.main import APP_DIR

    html = (APP_DIR / "index.html").read_text(encoding="utf-8")
    js = (APP_DIR / "app.js").read_text(encoding="utf-8")
    assert "promo-code" in html
    assert "promo_code" in js
