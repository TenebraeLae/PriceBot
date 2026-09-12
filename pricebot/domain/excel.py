from collections.abc import Mapping
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from pricebot.domain.catalog import CatalogProduct

COLUMN_ALIASES: dict[str, frozenset[str]] = {
    "sku": frozenset({"sku", "артикул", "код"}),
    "name": frozenset({"name", "название", "наименование"}),
    "category": frozenset({"category", "категория"}),
    "description": frozenset({"description", "описание"}),
    "price": frozenset({"price", "цена"}),
    "stock": frozenset({"stock", "остаток", "количество"}),
    "unit": frozenset({"unit", "единица", "ед"}),
    "photo": frozenset({"photo", "фото", "image", "картинка"}),
    "active": frozenset({"active", "активен"}),
    "sort": frozenset({"sort", "сортировка", "порядок"}),
}

REQUIRED_COLUMNS = ("sku", "name", "price")
EXPORT_COLUMNS = (
    "sku",
    "name",
    "category",
    "description",
    "price",
    "stock",
    "unit",
    "photo",
    "active",
    "sort",
)
NON_NUMERIC_PRICE_MARKERS = frozenset({"бесплатно", "free"})


@dataclass(frozen=True)
class RowError:
    row: int
    message: str


@dataclass
class ParseResult:
    rows_ok: list[CatalogProduct] = field(default_factory=list)
    rows_err: list[RowError] = field(default_factory=list)
    accepted: bool = True
    reject_reason: str | None = None

    @property
    def import_status(self) -> str:
        if not self.accepted:
            return "rejected"
        if not self.rows_ok:
            return "rejected"
        return "applied"


def normalize_header(value: Any) -> str:
    return str(value or "").strip().casefold().rstrip(".")


def _map_headers(headers: list[Any]) -> dict[str, int] | None:
    index_by_field: dict[str, int] = {}
    for idx, raw in enumerate(headers):
        name = normalize_header(raw)
        if not name:
            continue
        for field_name, aliases in COLUMN_ALIASES.items():
            if name in aliases and field_name not in index_by_field:
                index_by_field[field_name] = idx
    if any(name not in index_by_field for name in REQUIRED_COLUMNS):
        return None
    return index_by_field


def _cell(row: tuple[Any, ...], mapping: Mapping[str, int], key: str) -> Any:
    idx = mapping.get(key)
    if idx is None or idx >= len(row):
        return None
    return row[idx]


def _parse_decimal(value: Any) -> Decimal:
    if value is None or (isinstance(value, str) and not value.strip()):
        raise InvalidOperation
    if isinstance(value, Decimal):
        return value
    if isinstance(value, bool):
        raise InvalidOperation
    if isinstance(value, int):
        return Decimal(value)
    if isinstance(value, float):
        return Decimal(str(value))
    text = str(value).strip().replace(" ", "").replace(",", ".")
    if text.casefold() in NON_NUMERIC_PRICE_MARKERS:
        raise InvalidOperation
    return Decimal(text)


def _parse_active(value: Any) -> bool:
    if value is None or value == "":
        return True
    if isinstance(value, bool):
        return value
    text = str(value).strip().casefold()
    if text in {"1", "true", "да", "yes", "y"}:
        return True
    if text in {"0", "false", "нет", "no", "n"}:
        return False
    return True


def _parse_row(raw: tuple[Any, ...], mapping: Mapping[str, int]) -> CatalogProduct:
    sku = str(_cell(raw, mapping, "sku") or "").strip()
    name = str(_cell(raw, mapping, "name") or "").strip()
    if not sku or not name:
        raise ValueError("Нет sku или названия")

    price_raw = _cell(raw, mapping, "price")
    try:
        price = _parse_decimal(price_raw)
    except InvalidOperation as exc:
        raise ValueError("Нечисловая или пустая цена") from exc
    if price <= 0:
        raise ValueError("Цена должна быть больше 0")

    stock_raw = _cell(raw, mapping, "stock")
    if stock_raw is None or (isinstance(stock_raw, str) and not str(stock_raw).strip()):
        stock = Decimal("0")
    else:
        try:
            stock = _parse_decimal(stock_raw)
        except InvalidOperation as exc:
            raise ValueError("Некорректный остаток") from exc
        if stock < 0:
            raise ValueError("Отрицательный остаток")

    sort_raw = _cell(raw, mapping, "sort")
    sort = 0
    if sort_raw not in (None, ""):
        sort = int(_parse_decimal(sort_raw))

    unit_raw = _cell(raw, mapping, "unit")
    unit = str(unit_raw).strip() if unit_raw not in (None, "") else "шт"
    photo_raw = _cell(raw, mapping, "photo")
    photo = str(photo_raw).strip() if photo_raw not in (None, "") else None
    category = str(_cell(raw, mapping, "category") or "").strip()
    description = str(_cell(raw, mapping, "description") or "").strip()
    return CatalogProduct(
        sku=sku,
        name=name,
        category=category,
        description=description,
        price=price,
        stock=stock,
        unit=unit,
        photo_url=photo,
        active=_parse_active(_cell(raw, mapping, "active")),
        sort=sort,
    )


def parse_excel(
    source: bytes | Path | str,
    *,
    max_rows: int = 50_000,
) -> ParseResult:
    try:
        if isinstance(source, bytes):
            if not source:
                return ParseResult(accepted=False, reject_reason="empty_file")
            workbook = load_workbook(BytesIO(source), read_only=True, data_only=True)
        else:
            workbook = load_workbook(filename=str(source), read_only=True, data_only=True)
    except Exception:
        return ParseResult(accepted=False, reject_reason="unreadable_file")

    try:
        sheet = workbook.active
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return ParseResult(accepted=False, reject_reason="missing_headers")
        mapping = _map_headers(list(header_row or ()))
        if mapping is None:
            return ParseResult(accepted=False, reject_reason="missing_headers")

        result = ParseResult()
        seen_sku: set[str] = set()
        data_rows = 0
        for offset, raw in enumerate(rows_iter, start=2):
            if raw is None or all(cell is None or str(cell).strip() == "" for cell in raw):
                continue
            data_rows += 1
            if data_rows > max_rows:
                return ParseResult(accepted=False, reject_reason="too_many_rows")
            try:
                product = _parse_row(tuple(raw), mapping)
            except (ValueError, InvalidOperation, TypeError) as exc:
                result.rows_err.append(RowError(row=offset, message=str(exc)))
                continue
            if product.sku in seen_sku:
                result.rows_err.append(RowError(row=offset, message="Дубль sku"))
                continue
            seen_sku.add(product.sku)
            result.rows_ok.append(product)
        if not result.rows_ok:
            result.accepted = False
            if result.reject_reason is None:
                result.reject_reason = "no_valid_rows"
        return result
    finally:
        workbook.close()


def export_catalog_xlsx(products: list[CatalogProduct]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(EXPORT_COLUMNS))
    for product in products:
        sheet.append(
            [
                product.sku,
                product.name,
                product.category,
                product.description,
                str(product.price),
                str(product.stock),
                product.unit,
                product.photo_url or "",
                1 if product.active else 0,
                product.sort,
            ]
        )
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
